"""Downloadable analysis report (PDF) and data exports (CSV / Excel)."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from fpdf.fonts import FontFace

from audio_preprocessing.loader import load_mono

from ..utils import fmt_bytes, fmt_dt, fmt_pct
from . import storage, visuals

DISCLAIMER = ("SonicSentinel AI is a competition prototype for controlled and ethical testing. "
              "It is not a certified emergency-response or law-enforcement system.")

_REPLACE = {"–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"',
            "→": "->", "≥": ">=", "≤": "<=", "·": "-", "…": "...", "≈": "~"}


def latin(text) -> str:
    s = "" if text is None else str(text)
    for k, v in _REPLACE.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


class ReportPDF(FPDF):
    ink = (29, 28, 26)
    muted = (110, 106, 98)
    rule = (221, 217, 208)
    accent = (228, 87, 46)

    def header(self):
        self.set_fill_color(*self.accent)
        self.rect(15, 11, 2.2, 2.2, style="F")
        self.set_xy(19, 9.5)
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(*self.ink)
        self.cell(80, 5, "SonicSentinel AI", new_x=XPos.RIGHT, new_y=YPos.TOP)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*self.muted)
        self.cell(0, 5, latin(self.title_text), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_draw_color(*self.rule)
        self.line(15, 16, 195, 16)
        self.ln(6)

    def footer(self):
        self.set_y(-13)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(*self.muted)
        self.cell(150, 4, latin(DISCLAIMER))
        self.cell(0, 4, f"Page {self.page_no()}", align="R")

    def section(self, title: str):
        self.ln(3)
        self.set_font("Helvetica", "B", 10)
        self.set_text_color(*self.ink)
        self.cell(0, 6, latin(title.upper()), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_draw_color(*self.rule)
        self.line(15, self.get_y(), 195, self.get_y())
        self.ln(1.5)

    def kv_table(self, rows: list[tuple[str, str]], widths=(55, 125)):
        self.set_font("Helvetica", "", 8.5)
        self.set_text_color(*self.ink)
        with self.table(col_widths=widths, borders_layout="HORIZONTAL_LINES", line_height=4.8,
                        first_row_as_headings=False, text_align=("LEFT", "LEFT")) as table:
            for k, v in rows:
                r = table.row()
                r.cell(latin(k))
                r.cell(latin(v))

    def grid(self, headings: list[str], rows: list[list], widths=None, font_size: float = 7.5):
        self.set_font("Helvetica", "", font_size)
        self.set_text_color(*self.ink)
        heading = FontFace(emphasis="BOLD", color=self.muted, fill_color=(248, 246, 241))
        with self.table(col_widths=widths, borders_layout="HORIZONTAL_LINES", line_height=4.3,
                        headings_style=heading) as table:
            h = table.row()
            for x in headings:
                h.cell(latin(x))
            for row in rows:
                r = table.row()
                for x in row:
                    r.cell(latin(x))


def _pct(v) -> str:
    return fmt_pct(v, 1) if v is not None else "-"


def record_pdf(record) -> bytes:
    pdf = ReportPDF(orientation="P", unit="mm", format="A4")
    pdf.title_text = f"Sound event analysis report - {record.audio_id}"
    pdf.set_title(f"SonicSentinel report {record.audio_id}")
    pdf.set_author("SonicSentinel AI")
    pdf.set_auto_page_break(True, 16)
    pdf.set_margins(15, 12, 15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(*pdf.ink)
    pdf.cell(0, 9, latin(f"{record.effective_category or 'Pending'}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*pdf.muted)
    pdf.cell(0, 5, latin(f"{record.audio_id}  -  {record.display_name}  -  generated {datetime.utcnow():%d %b %Y %H:%M} UTC"),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    d = record.decision or {}
    pdf.section("Final decision")
    latest_review = record.reviews[-1] if record.reviews else None
    pdf.kv_table([
        ("Detected sound category", record.final_category or "-"),
        ("Confidence (fused)", f"{_pct(record.final_confidence)} ({record.confidence_level or '-'})"),
        ("Model agreement status", record.consistency or "-"),
        ("Event severity", record.effective_severity or "-"),
        ("Alert status", record.alert_status or "-"),
        ("Manual review required", "Yes - " + "; ".join(record.review_reasons or []) if record.needs_review else "No"),
        ("Recommended action", record.recommended_action or "-"),
        ("Event status", record.status or "-"),
        ("Reviewer decision", f"{latest_review.decision}: {latest_review.original_category} -> "
                              f"{latest_review.reviewed_category} by {latest_review.reviewer.username}"
                              + (f" ({latest_review.comment})" if latest_review.comment else "") if latest_review else "Not reviewed"),
    ])

    pdf.section("Audio metadata")
    pdf.kv_table([
        ("Filename", record.original_filename or "-"),
        ("Source", record.source),
        ("Format / codec", f"{record.format or '-'} / {record.codec or '-'}"),
        ("Duration", f"{record.duration or 0:.2f} s"),
        ("Sampling rate / channels / bit depth",
         f"{record.sample_rate or '-'} Hz / {record.channels or '-'} / {record.bit_depth or 'n/a'}"),
        ("File size", fmt_bytes(record.file_size)),
        ("Uploaded", f"{fmt_dt(record.created_at)} by {record.user.username if record.user else '-'}"),
        ("SHA-256", record.sha256 or "-"),
        ("Duplicate check", (f"identical to {record.duplicate_of.audio_id}" if record.duplicate_of else
                             f"near-duplicate of {record.near_duplicate_of.audio_id} (BER {record.near_duplicate_ber})"
                             if record.near_duplicate_of else "no duplicate found")),
        ("Model versions", f"Python {record.python_model_version or '-'}; GTM {record.gtm_model_version or 'n/a'}"),
    ])

    pdf.section("Python model vs Google Teachable Machine")
    py, gtm = d.get("python") or {}, d.get("gtm") or {}
    pdf.kv_table([
        ("Python prediction", f"{record.python_top or '-'} ({_pct(record.python_confidence)}), top-two margin "
                              f"{py.get('margin', 0):.2f}"),
        ("GTM prediction", f"{record.gtm_top or 'not available'} ({_pct(record.gtm_confidence)})"
                           + (f", top-two margin {gtm.get('margin', 0):.2f}" if gtm else "")),
        ("Predicted-class match", "-" if record.class_match is None else ("Yes" if record.class_match else "No")),
        ("Top-class confidence difference", "-" if record.confidence_difference is None else f"{record.confidence_difference:.3f}"),
        ("Top-two margin difference", "-" if record.margin_difference is None else f"{record.margin_difference:.3f}"),
    ])
    classes = list((record.python_scores or {}).keys())
    fused = d.get("fused_scores") or {}
    rows = [[c, _pct((record.python_scores or {}).get(c)), _pct((record.gtm_scores or {}).get(c)) if record.gtm_scores else "-",
             _pct(fused.get(c))] for c in classes]
    rows.sort(key=lambda r: -(fused.get(r[0]) or 0))
    pdf.ln(1)
    pdf.grid(["Class", "Python confidence", "GTM confidence", "Fused"], rows, widths=(70, 36, 36, 38))

    pdf.section("Audio quality")
    q = record.quality or {}
    m = q.get("metrics", {})
    pdf.kv_table([
        ("Quality", f"{record.quality_label or '-'} (score {record.quality_score})"),
        ("Levels", f"peak {m.get('peak_dbfs')} dBFS, RMS {m.get('rms_dbfs')} dBFS, noise floor {m.get('noise_floor_dbfs')} dBFS"),
        ("Estimated SNR / silence / clipping", f"{m.get('snr_estimate_db')} dB / {fmt_pct(m.get('silence_ratio'), 0)} / "
                                               f"{fmt_pct(m.get('clipping_ratio'), 2)}"),
        ("Issues", "; ".join(f"{i['label']}: {i['detail']}" for i in q.get("issues", [])) or "none"),
    ])

    src = storage.resolve(record.stored_path)
    if src is not None:
        try:
            y, sr = load_mono(src)
            wave_png, spec_png = visuals.report_figure(y, sr, record.events or [])
            pdf.section("Waveform")
            pdf.image(io.BytesIO(wave_png), w=180)
            pdf.section("Mel spectrogram")
            pdf.image(io.BytesIO(spec_png), w=180)
        except Exception as exc:  # report must still be produced without figures
            pdf.section("Waveform and spectrogram")
            pdf.multi_cell(0, 5, latin(f"Figures unavailable: {exc}"))
    else:
        pdf.section("Waveform and spectrogram")
        pdf.multi_cell(0, 5, "Audio file not stored (privacy setting) or removed by the retention policy.")

    if record.segments and len(record.segments) > 1:
        pdf.section("Segments")
        seg_rows = []
        for s in record.segments[:40]:
            seg_rows.append([s.index, f"{s.start_time:.1f}-{s.end_time:.1f}s",
                             "silent" if s.silent else f"{s.python_top} {_pct(s.python_confidence)}",
                             f"{s.gtm_top} {_pct(s.gtm_confidence)}" if s.gtm_top else "-",
                             s.final_category or "-", s.severity or "-", s.alert_status or "-"])
        pdf.grid(["#", "Time", "Python", "GTM", "Final", "Severity", "Alert"], seg_rows,
                 widths=(8, 22, 36, 36, 32, 20, 26), font_size=7)

    if record.alerts:
        pdf.section("Alerts")
        pdf.grid(["Alert", "Category", "Severity", "Status", "Created", "Actions"],
                 [[a.alert_code, a.category, a.severity, a.status, fmt_dt(a.created_at),
                   "; ".join(f"{x.action} by {x.user.username if x.user else '-'}" for x in a.actions) or "-"]
                  for a in record.alerts], widths=(28, 32, 20, 22, 34, 44))

    pdf.section("Pre-processing")
    steps = (record.preprocessing or {}).get("steps", [])
    pdf.grid(["Step", "Detail"], [[s["step"], s["detail"]] for s in steps], widths=(50, 130))
    return bytes(pdf.output())


# --------------------------------------------------------------------------- #
# CSV / Excel export
# --------------------------------------------------------------------------- #
def record_rows(records, classes: list[str]) -> tuple[list[str], list[list]]:
    head = ["audio_id", "created_at_utc", "source", "filename", "user", "duration_s", "format", "sample_rate",
            "channels", "quality", "quality_score", "python_prediction", "python_confidence", "gtm_prediction",
            "gtm_confidence", "class_match", "confidence_difference", "consistency", "final_category",
            "final_confidence", "confidence_level", "severity", "alert_status", "status", "needs_review",
            "review_reasons", "reviewed_category", "reviewer", "python_model_version", "gtm_model_version",
            "duplicate_of", "sha256"]
    head += [f"python_{c}" for c in classes] + [f"gtm_{c}" for c in classes]
    rows = []
    for r in records:
        ps, gs = r.python_scores or {}, r.gtm_scores or {}
        rows.append([
            r.audio_id, r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "", r.source,
            r.original_filename, r.user.username if r.user else "", r.duration, r.format, r.sample_rate, r.channels,
            r.quality_label, r.quality_score, r.python_top, r.python_confidence, r.gtm_top, r.gtm_confidence,
            r.class_match, r.confidence_difference, r.consistency, r.final_category, r.final_confidence,
            r.confidence_level, r.effective_severity, r.alert_status, r.status, r.needs_review,
            "; ".join(r.review_reasons or []), r.reviewed_category, r.reviewer.username if r.reviewer else "",
            r.python_model_version, r.gtm_model_version, r.duplicate_of.audio_id if r.duplicate_of else "", r.sha256,
        ] + [ps.get(c) for c in classes] + [gs.get(c) for c in classes])
    return head, rows


def to_csv(head: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(head)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, head, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(head)
        for row in rows:
            ws.append([("" if v is None else v) for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D1C1A")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(head, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(9, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
